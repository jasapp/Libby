"""Merge results/*.jsonl into a copy of the source workbook, filling column L (Website)."""
import openpyxl, json, glob, sys
from openpyxl.styles import Font
SRC='data/US_Public_Libraries_10k_plus.xlsx'
OUT=sys.argv[1] if len(sys.argv)>1 else 'US_Public_Libraries_10k_plus_with_websites.xlsx'
res={}
for f in sorted(glob.glob('data/results/*.jsonl')):
    for line in open(f):
        line=line.strip()
        if not line: continue
        try: o=json.loads(line)
        except Exception as e: print('bad line',f,line[:80]); continue
        res[o['id']]=o
wb=openpyxl.load_workbook(SRC)
ws=wb['Libraries 10k+']
hdr=[c.value for c in ws[1]]
col_site=hdr.index('Website (fill in)')+1
col_id=hdr.index('IMLS ID')+1
# add confidence column at end
col_conf=ws.max_column+1
ws.cell(1,col_conf,'Website Confidence').font=Font(bold=True)
filled=0; blank=0; missing=0
for r in range(2, ws.max_row+1):
    lid=ws.cell(r,col_id).value
    o=res.get(lid)
    if not o: missing+=1; continue
    url=(o.get('website') or '').strip()
    if url:
        c=ws.cell(r,col_site,url); c.hyperlink=url; c.font=Font(color='0563C1',underline='single')
        filled+=1
    else: blank+=1
    ws.cell(r,col_conf,o.get('confidence',''))
ws.cell(1,col_site).value='Website'
notes=wb['Notes']
notes.cell(notes.max_row+1,1,f'Website column populated via web search per library ({filled} found, {blank} not found). Website Confidence: high = clearly the official site; medium = probably right; low = unsure. Verify low-confidence rows before use.')
wb.save(OUT)
print(dict(filled=filled, blank=blank, missing=missing, total=ws.max_row-1))
